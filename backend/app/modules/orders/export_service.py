"""Export service for SmartDine order, billing, and culinary ticket histories.
Supports Excel (.xlsx) and PDF (landscape) with role-based masking,
Karachi timezone alignment, and Excel formula-injection safety.
"""
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
import io
from typing import Any, Literal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.exceptions import AppError
from app.core.security import Actor

KARACHI_TZ = ZoneInfo("Asia/Karachi")


def karachi_date_range_to_utc(start_date: date, end_date: date) -> tuple[datetime, datetime]:
    """Convert inclusive calendar dates in Asia/Karachi to UTC datetime boundaries."""
    if end_date < start_date:
        raise AppError("invalid_date_range", "End date must be on or after start date.", 422)
    if (end_date - start_date).days > 366:
        raise AppError("invalid_date_range", "Select a reporting period of at most 366 days.", 422)

    start_local = datetime.combine(start_date, time.min, tzinfo=KARACHI_TZ)
    end_local = datetime.combine(end_date + timedelta(days=1), time.min, tzinfo=KARACHI_TZ)
    return start_local.astimezone(timezone.utc), end_local.astimezone(timezone.utc)


def sanitize_excel_cell(val: Any) -> Any:
    """Neutralize potential CSV/Excel formula injection (e.g. =, +, -, @)."""
    if isinstance(val, str) and val.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{val}"
    return val


def format_money_str(val: Any) -> str:
    try:
        num = float(val or 0)
        return f"PKR {num:,.2f}"
    except (ValueError, TypeError):
        return "PKR 0.00"


def format_karachi_dt(dt: datetime | None, include_time: bool = True) -> str:
    if not dt:
        return "—"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    local_dt = dt.astimezone(KARACHI_TZ)
    if include_time:
        return local_dt.strftime("%d %b %Y, %I:%M %p")
    return local_dt.strftime("%d %b %Y")


async def fetch_orders_for_export(
    gateway: Any,
    actor: Actor,
    start_date: date,
    end_date: date,
    status: str | None = None,
    q: str = "",
    scope: Literal["manager", "cashier", "waiter", "kitchen"] = "manager",
) -> tuple[list[dict], dict]:
    """Fetch orders and their line items scoped to branch, role and date range."""
    start_utc, end_utc = karachi_date_range_to_utc(start_date, end_date)

    # 1. Resolve branch_id and branch name from actor
    prof_row = await gateway.query_one(
        "SELECT p.branch_id, b.name as branch_name, p.full_name FROM private.profiles p "
        "LEFT JOIN private.branches b ON b.id = p.branch_id WHERE p.id = %s",
        (str(actor.id),),
    )
    branch_id = prof_row["branch_id"] if prof_row else None
    branch_name = prof_row["branch_name"] if prof_row and prof_row.get("branch_name") else "SmartDine"

    # 2. Build parameterized SQL query
    conditions = ["o.created_at >= %s", "o.created_at < %s"]
    params: list[Any] = [start_utc, end_utc]

    if branch_id:
        conditions.append("o.branch_id = %s")
        params.append(branch_id)

    # Scoping rules
    if scope == "waiter" or (actor.role != "manager" and actor.staff_type == "waiter"):
        conditions.append("o.created_by = %s")
        params.append(str(actor.id))

    if status and status.lower() != "all":
        conditions.append("o.status = %s")
        params.append(status.lower())

    if q.strip():
        search_pattern = f"%{q.strip().lower()}%"
        conditions.append(
            """(
                LOWER(COALESCE(o.order_number, '')) LIKE %s
                OR LOWER(COALESCE(o.table_name_snapshot, '')) LIKE %s
                OR LOWER(COALESCE(o.floor_name_snapshot, '')) LIKE %s
                OR LOWER(COALESCE(o.notes, '')) LIKE %s
                OR LOWER(COALESCE(o.created_by_name, '')) LIKE %s
                OR LOWER(COALESCE(o.prepared_by_name, '')) LIKE %s
                OR LOWER(COALESCE(o.paid_by_name, '')) LIKE %s
                OR EXISTS (
                    SELECT 1 FROM private.order_items oi
                    WHERE oi.order_id = o.id AND LOWER(oi.name_snapshot) LIKE %s
                )
            )"""
        )
        params.extend([search_pattern] * 8)

    where_clause = " AND ".join(conditions)

    orders_sql = f"""
        SELECT
            o.id,
            o.order_number,
            o.branch_id,
            o.status,
            o.created_at,
            o.updated_at,
            o.prepared_at,
            o.completed_at,
            o.cancelled_at,
            o.cancellation_reason,
            o.floor_name_snapshot,
            o.table_name_snapshot,
            o.seats_snapshot,
            o.created_by,
            o.created_by_name,
            o.created_by_email,
            o.prepared_by,
            o.prepared_by_name,
            o.prepared_by_email,
            o.paid_by,
            o.paid_by_name,
            o.paid_by_email,
            o.notes,
            o.tax,
            o.tax_rate_snapshot,
            o.discount,
            o.discount_percent,
            o.discount_reason,
            o.cash_received,
            o.change_given
        FROM private.orders o
        WHERE {where_clause}
        ORDER BY o.created_at DESC
    """
    order_rows = await gateway.query(orders_sql, tuple(params))
    if not order_rows:
        return [], {"branch_name": branch_name}

    order_ids = [r["id"] for r in order_rows]

    # Fetch order items
    items_sql = """
        SELECT
            order_id,
            menu_item_id,
            name_snapshot,
            quantity,
            price_snapshot
        FROM private.order_items
        WHERE order_id = ANY(%s)
        ORDER BY id ASC
    """
    item_rows = await gateway.query(items_sql, (order_ids,))

    items_by_order: dict[Any, list[dict]] = {}
    for item in item_rows:
        items_by_order.setdefault(item["order_id"], []).append(item)

    normalized_orders = []
    for o in order_rows:
        items = items_by_order.get(o["id"], [])
        subtotal = sum((Decimal(str(it["price_snapshot"] or 0)) * Decimal(it["quantity"]) for it in items), Decimal("0.00"))
        discount = Decimal(str(o.get("discount") or 0))
        tax = Decimal(str(o.get("tax") or 0))
        total = max(Decimal("0.00"), subtotal - discount + tax)

        # Calculate preparation duration in minutes if prepared_at is available
        prep_minutes = None
        if o.get("prepared_at") and o.get("created_at"):
            diff = (o["prepared_at"] - o["created_at"]).total_seconds()
            prep_minutes = max(0, round(diff / 60))
        elif o.get("completed_at") and o.get("created_at"):
            diff = (o["completed_at"] - o["created_at"]).total_seconds()
            prep_minutes = max(0, round(diff / 60))

        norm = {
            "id": str(o["id"]),
            "order_number": o.get("order_number") or f"ORD-{str(o['id'])[:8]}",
            "receipt_number": f"RCP-{o.get('order_number') or str(o['id'])[:8]}",
            "status": o.get("status") or "pending",
            "created_at": o.get("created_at"),
            "prepared_at": o.get("prepared_at"),
            "completed_at": o.get("completed_at"),
            "cancelled_at": o.get("cancelled_at"),
            "cancellation_reason": o.get("cancellation_reason"),
            "prep_minutes": prep_minutes,
            "floor_table": " · ".join(filter(None, [
                f"Floor {o['floor_name_snapshot']}" if o.get("floor_name_snapshot") else "",
                o.get("table_name_snapshot") or "Dining room",
            ])),
            "table_name": o.get("table_name_snapshot") or "Dining room",
            "seats": o.get("seats_snapshot"),
            "waiter_name": o.get("created_by_name") or "—",
            "chef_name": o.get("prepared_by_name") or "—",
            "cashier_name": o.get("paid_by_name") or "—",
            "notes": o.get("notes") or "",
            "subtotal": subtotal,
            "discount": discount,
            "discount_percent": o.get("discount_percent"),
            "discount_reason": o.get("discount_reason"),
            "tax": tax,
            "tax_rate": o.get("tax_rate_snapshot") or Decimal("15.00"),
            "total": total,
            "cash_received": Decimal(str(o.get("cash_received") or 0)) if o.get("cash_received") is not None else None,
            "change_given": Decimal(str(o.get("change_given") or 0)) if o.get("change_given") is not None else None,
            "items": [
                {
                    "name": it.get("name_snapshot") or "Item",
                    "quantity": int(it.get("quantity") or 1),
                    "unit_price": Decimal(str(it.get("price_snapshot") or 0)),
                    "line_total": Decimal(str(it.get("price_snapshot") or 0)) * Decimal(it.get("quantity") or 1),
                }
                for it in items
            ],
            "items_summary": ", ".join(f"{it.get('quantity')}x {it.get('name_snapshot')}" for it in items),
        }
        normalized_orders.append(norm)

    meta = {
        "branch_name": branch_name,
        "total_records": len(normalized_orders),
    }
    return normalized_orders, meta


# ============================================================================
# EXCEL WORKBOOK BUILDER (openpyxl)
# ============================================================================

def build_excel_export(
    records: list[dict],
    scope: Literal["manager", "cashier", "waiter", "kitchen"],
    start_date: date,
    end_date: date,
    actor: Actor,
    meta: dict,
) -> io.BytesIO:
    wb = Workbook()

    # Brand Styles
    brand_font = Font(name="Calibri", size=11)
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    section_font = Font(name="Calibri", size=12, bold=True, color="1E293B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    kpi_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    pkr_currency_fmt = '_("PKR"* #,##0.00_);_("PKR"* (#,##0.00);_("PKR"* "-"??_);_(@_)'

    # -------------------------------------------------------------------------
    # SHEET 1: SUMMARY
    # -------------------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary["A1"] = "SmartDine AI — Operational History Report"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = (
        f"Branch: {meta.get('branch_name', 'SmartDine')} | Scope: {scope.capitalize()} Portal | "
        f"Period: {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')} (Asia/Karachi, Inclusive)"
    )
    ws_summary["A2"].font = subtitle_font

    ws_summary["A4"] = "Report Metadata"
    ws_summary["A4"].font = section_font
    meta_rows = [
        ("Generated By", f"{actor.full_name} ({actor.role})"),
        ("Generated Time", format_karachi_dt(datetime.now(timezone.utc))),
        ("Total Records Exported", len(records)),
        ("Date Filter (Karachi)", f"{start_date.isoformat()} to {end_date.isoformat()}"),
        ("Portal View Scoping", f"{scope.upper()} operational projection"),
    ]
    r_idx = 5
    for label, val in meta_rows:
        ws_summary.cell(row=r_idx, column=1, value=label).font = bold_font
        ws_summary.cell(row=r_idx, column=2, value=sanitize_excel_cell(val)).font = brand_font
        ws_summary.cell(row=r_idx, column=1).border = thin_border
        ws_summary.cell(row=r_idx, column=2).border = thin_border
        r_idx += 1

    r_idx += 1
    ws_summary.cell(row=r_idx, column=1, value="Key Performance Summary").font = section_font
    r_idx += 1

    completed_count = sum(1 for r in records if r["status"] == "completed")
    cancelled_count = sum(1 for r in records if r["status"] == "cancelled")
    active_count = len(records) - completed_count - cancelled_count

    kpis: list[tuple[str, Any, str]] = [
        ("Total Orders / Tickets", len(records), "int"),
        ("Completed / Served", completed_count, "int"),
        ("Cancelled", cancelled_count, "int"),
        ("Pending / In Progress", active_count, "int"),
    ]

    if scope in ("manager", "cashier"):
        total_subtotal = sum(r["subtotal"] for r in records)
        total_discount = sum(r["discount"] for r in records)
        total_tax = sum(r["tax"] for r in records)
        total_net = sum(r["total"] for r in records)
        kpis.extend([
            ("Gross Subtotal", float(total_subtotal), "money"),
            ("Total Discounts Given", float(total_discount), "money"),
            ("Total Tax Collected", float(total_tax), "money"),
            ("Net Settled Sales", float(total_net), "money"),
        ])
    elif scope == "kitchen":
        durations = [r["prep_minutes"] for r in records if r.get("prep_minutes") is not None]
        avg_prep = round(sum(durations) / len(durations), 1) if durations else 0
        total_dishes = sum(sum(it["quantity"] for it in r["items"]) for r in records)
        kpis.extend([
            ("Total Dishes Dispatched", total_dishes, "int"),
            ("Average Prep Time", f"{avg_prep} minutes", "str"),
        ])

    for label, val, kind in kpis:
        cell_lbl = ws_summary.cell(row=r_idx, column=1, value=label)
        cell_val = ws_summary.cell(row=r_idx, column=2, value=val if kind != "str" else sanitize_excel_cell(val))
        cell_lbl.font = bold_font
        cell_val.font = bold_font
        cell_lbl.border = thin_border
        cell_val.border = thin_border
        cell_lbl.fill = kpi_fill
        cell_val.fill = kpi_fill
        if kind == "money":
            cell_val.number_format = pkr_currency_fmt
            cell_val.alignment = align_right
        elif kind == "int":
            cell_val.number_format = "#,##0"
            cell_val.alignment = align_right
        else:
            cell_val.alignment = align_left
        r_idx += 1

    # Daily aggregation table
    r_idx += 2
    ws_summary.cell(row=r_idx, column=1, value="Daily Activity Breakdown").font = section_font
    r_idx += 1
    day_headers = ["Date", "Orders", "Completed", "Cancelled"]
    if scope in ("manager", "cashier"):
        day_headers.append("Net Revenue (PKR)")

    for c_i, h_text in enumerate(day_headers, start=1):
        cell = ws_summary.cell(row=r_idx, column=c_i, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
    r_idx += 1

    daily_map: dict[str, dict[str, Any]] = {}
    for r in records:
        day_str = format_karachi_dt(r["created_at"], include_time=False)
        daily_map.setdefault(day_str, {"count": 0, "completed": 0, "cancelled": 0, "revenue": Decimal("0.00")})
        daily_map[day_str]["count"] += 1
        if r["status"] == "completed":
            daily_map[day_str]["completed"] += 1
        elif r["status"] == "cancelled":
            daily_map[day_str]["cancelled"] += 1
        daily_map[day_str]["revenue"] += r["total"]

    for day_str, stats in sorted(daily_map.items()):
        ws_summary.cell(row=r_idx, column=1, value=day_str).alignment = align_center
        ws_summary.cell(row=r_idx, column=2, value=stats["count"]).alignment = align_right
        ws_summary.cell(row=r_idx, column=3, value=stats["completed"]).alignment = align_right
        ws_summary.cell(row=r_idx, column=4, value=stats["cancelled"]).alignment = align_right
        for c in range(1, 5):
            ws_summary.cell(row=r_idx, column=c).border = thin_border
            ws_summary.cell(row=r_idx, column=c).font = brand_font
        if scope in ("manager", "cashier"):
            rev_cell = ws_summary.cell(row=r_idx, column=5, value=float(stats["revenue"]))
            rev_cell.number_format = pkr_currency_fmt
            rev_cell.alignment = align_right
            rev_cell.border = thin_border
            rev_cell.font = brand_font
        r_idx += 1

    # -------------------------------------------------------------------------
    # SHEET 2: HISTORY RECORDS
    # -------------------------------------------------------------------------
    ws_hist = wb.create_sheet(title="History")
    ws_hist.views.sheetView[0].showGridLines = True
    ws_hist.freeze_panes = "A2"

    columns_by_scope = {
        "manager": [
            ("Order #", 14, align_left, "str"),
            ("Date & Time", 20, align_center, "str"),
            ("Floor & Table", 18, align_left, "str"),
            ("Status", 14, align_center, "str"),
            ("Waiter", 16, align_left, "str"),
            ("Kitchen Chef", 16, align_left, "str"),
            ("Cashier", 16, align_left, "str"),
            ("Items Ordered", 36, align_left, "str"),
            ("Subtotal", 14, align_right, "money"),
            ("Discount", 12, align_right, "money"),
            ("Tax", 12, align_right, "money"),
            ("Total (PKR)", 15, align_right, "money"),
            ("Special Notes", 24, align_left, "str"),
        ],
        "cashier": [
            ("Receipt #", 16, align_left, "str"),
            ("Order #", 14, align_left, "str"),
            ("Date & Time", 20, align_center, "str"),
            ("Table", 16, align_left, "str"),
            ("Status", 14, align_center, "str"),
            ("Subtotal", 14, align_right, "money"),
            ("Discount %", 12, align_right, "str"),
            ("Discount", 12, align_right, "money"),
            ("Tax", 12, align_right, "money"),
            ("Total (PKR)", 15, align_right, "money"),
            ("Cash Received", 15, align_right, "money"),
            ("Change Given", 14, align_right, "money"),
            ("Settled By", 16, align_left, "str"),
            ("Notes", 22, align_left, "str"),
        ],
        "waiter": [
            ("Order #", 14, align_left, "str"),
            ("Date & Time", 20, align_center, "str"),
            ("Floor & Table", 18, align_left, "str"),
            ("Status", 14, align_center, "str"),
            ("Items Ordered", 38, align_left, "str"),
            ("Bill Total", 15, align_right, "money"),
            ("Special Notes", 28, align_left, "str"),
        ],
        "kitchen": [
            ("Ticket #", 14, align_left, "str"),
            ("Order #", 14, align_left, "str"),
            ("Order Time", 20, align_center, "str"),
            ("Prepared Time", 20, align_center, "str"),
            ("Prep Duration", 16, align_center, "str"),
            ("Floor & Table", 18, align_left, "str"),
            ("Status", 14, align_center, "str"),
            ("Dishes & Quantities", 42, align_left, "str"),
            ("Chef Dispatcher", 16, align_left, "str"),
            ("Chef Notes", 28, align_left, "str"),
        ],
    }

    cols = columns_by_scope.get(scope, columns_by_scope["manager"])

    for col_idx, (col_name, _, _, _) in enumerate(cols, start=1):
        cell = ws_hist.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    row_num = 2
    for r in records:
        if scope == "manager":
            row_data = [
                (r["order_number"], "str"),
                (format_karachi_dt(r["created_at"]), "str"),
                (r["floor_table"], "str"),
                (r["status"].upper(), "str"),
                (r["waiter_name"], "str"),
                (r["chef_name"], "str"),
                (r["cashier_name"], "str"),
                (r["items_summary"], "str"),
                (float(r["subtotal"]), "money"),
                (float(r["discount"]), "money"),
                (float(r["tax"]), "money"),
                (float(r["total"]), "money"),
                (r["notes"], "str"),
            ]
        elif scope == "cashier":
            disc_pct_str = f"{r['discount_percent']}%" if r.get("discount_percent") else "0%"
            row_data = [
                (r["receipt_number"], "str"),
                (r["order_number"], "str"),
                (format_karachi_dt(r["created_at"]), "str"),
                (r["floor_table"], "str"),
                (r["status"].upper(), "str"),
                (float(r["subtotal"]), "money"),
                (disc_pct_str, "str"),
                (float(r["discount"]), "money"),
                (float(r["tax"]), "money"),
                (float(r["total"]), "money"),
                (float(r["cash_received"]) if r.get("cash_received") is not None else 0.0, "money"),
                (float(r["change_given"]) if r.get("change_given") is not None else 0.0, "money"),
                (r["cashier_name"], "str"),
                (r["notes"], "str"),
            ]
        elif scope == "waiter":
            row_data = [
                (r["order_number"], "str"),
                (format_karachi_dt(r["created_at"]), "str"),
                (r["floor_table"], "str"),
                (r["status"].upper(), "str"),
                (r["items_summary"], "str"),
                (float(r["total"]), "money"),
                (r["notes"], "str"),
            ]
        else:  # kitchen
            prep_str = f"{r['prep_minutes']} min" if r.get("prep_minutes") is not None else "—"
            row_data = [
                (f"TCK-{r['order_number']}", "str"),
                (r["order_number"], "str"),
                (format_karachi_dt(r["created_at"]), "str"),
                (format_karachi_dt(r["prepared_at"]) if r.get("prepared_at") else "—", "str"),
                (prep_str, "str"),
                (r["floor_table"], "str"),
                (r["status"].upper(), "str"),
                (r["items_summary"], "str"),
                (r["chef_name"], "str"),
                (r["notes"], "str"),
            ]

        for col_idx, ((val, kind), (_, _, align, _)) in enumerate(zip(row_data, cols), start=1):
            cell = ws_hist.cell(row=row_num, column=col_idx)
            cell.font = brand_font
            cell.alignment = align
            cell.border = thin_border
            if kind == "money":
                cell.value = val
                cell.number_format = pkr_currency_fmt
            else:
                cell.value = sanitize_excel_cell(val)

        row_num += 1

    # Enable Excel Auto-Filter on History
    ws_hist.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{row_num - 1}"

    # -------------------------------------------------------------------------
    # SHEET 3: ITEMS BREAKDOWN (FOR PIVOT TABLES & DISH ANALYSIS)
    # -------------------------------------------------------------------------
    ws_items = wb.create_sheet(title="Items")
    ws_items.views.sheetView[0].showGridLines = True
    ws_items.freeze_panes = "A2"

    item_headers = ["Order #", "Date", "Table", "Dish Name", "Quantity"]
    if scope in ("manager", "cashier"):
        item_headers.extend(["Unit Price (PKR)", "Line Total (PKR)"])

    for c_i, h_text in enumerate(item_headers, start=1):
        cell = ws_items.cell(row=1, column=c_i, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    item_row_num = 2
    for r in records:
        for it in r["items"]:
            ws_items.cell(row=item_row_num, column=1, value=sanitize_excel_cell(r["order_number"])).alignment = align_left
            ws_items.cell(row=item_row_num, column=2, value=format_karachi_dt(r["created_at"], include_time=False)).alignment = align_center
            ws_items.cell(row=item_row_num, column=3, value=sanitize_excel_cell(r["floor_table"])).alignment = align_left
            ws_items.cell(row=item_row_num, column=4, value=sanitize_excel_cell(it["name"])).alignment = align_left
            qty_cell = ws_items.cell(row=item_row_num, column=5, value=it["quantity"])
            qty_cell.alignment = align_right
            qty_cell.number_format = "#,##0"

            for c in range(1, 6):
                ws_items.cell(row=item_row_num, column=c).border = thin_border
                ws_items.cell(row=item_row_num, column=c).font = brand_font

            if scope in ("manager", "cashier"):
                p_cell = ws_items.cell(row=item_row_num, column=6, value=float(it["unit_price"]))
                p_cell.number_format = pkr_currency_fmt
                p_cell.alignment = align_right
                p_cell.border = thin_border
                p_cell.font = brand_font

                tot_cell = ws_items.cell(row=item_row_num, column=7, value=float(it["line_total"]))
                tot_cell.number_format = pkr_currency_fmt
                tot_cell.alignment = align_right
                tot_cell.border = thin_border
                tot_cell.font = brand_font

            item_row_num += 1

    ws_items.auto_filter.ref = f"A1:{get_column_letter(len(item_headers))}{item_row_num - 1}"

    # Auto-set column widths across all sheets
    for ws in [ws_summary, ws_hist, ws_items]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = max(val.split("\n"), key=len)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ============================================================================
# PDF BUILDER (reportlab - Landscape with NumberedCanvas)
# ============================================================================

class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and render total page count 'Page X of Y'."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states: list[dict] = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_footer(num_pages)
            super().showPage()
        super().save()

    def draw_footer(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        # Top border line for footer
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 30, 806, 30)

        # Left footer: Branding & timestamp
        gen_time = format_karachi_dt(datetime.now(timezone.utc))
        self.drawString(36, 18, f"SmartDine AI Operations — Generated {gen_time} (PKT)")

        # Right footer: Page numbers
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(806, 18, page_str)
        self.restoreState()


def build_pdf_export(
    records: list[dict],
    scope: Literal["manager", "cashier", "waiter", "kitchen"],
    start_date: date,
    end_date: date,
    actor: Actor,
    meta: dict,
) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=42,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    style_title = ParagraphStyle(
        "ReportTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0F172A"),
    )
    style_meta_right = ParagraphStyle(
        "ReportMetaRight",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        alignment=2,  # Right-aligned
        textColor=colors.HexColor("#475569"),
    )
    style_th = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1,  # Center
    )
    style_td = ParagraphStyle(
        "TableBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1E293B"),
    )
    style_td_center = ParagraphStyle(
        "TableBodyCenter",
        parent=style_td,
        alignment=1,
    )
    style_td_right = ParagraphStyle(
        "TableBodyRight",
        parent=style_td,
        alignment=2,
    )
    style_kpi_num = ParagraphStyle(
        "KpiNum",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=15,
        textColor=colors.HexColor("#0F172A"),
        alignment=1,
    )
    style_kpi_lbl = ParagraphStyle(
        "KpiLbl",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9,
        textColor=colors.HexColor("#64748B"),
        alignment=1,
    )

    story = []

    # 1. Header Banner
    title_text = {
        "manager": "Executive Operations & Order History",
        "cashier": "Cashier & Billing Settlement History",
        "waiter": "Waiter Floor Order History",
        "kitchen": "Kitchen Culinary Dispatch History",
    }.get(scope, "Order History")

    header_table_data = [
        [
            Paragraph(
                f"<b>SMARTDINE AI</b> · {meta.get('branch_name', 'SmartDine')}<br/>"
                f"<font size=13 color='#0F172A'><b>{title_text}</b></font>",
                style_title,
            ),
            Paragraph(
                f"<b>Period:</b> {start_date.strftime('%d %b %Y')} – {end_date.strftime('%d %b %Y')}<br/>"
                f"<b>Records:</b> {len(records)} orders | <b>Operator:</b> {actor.full_name} ({actor.role})<br/>"
                f"<b>Timezone:</b> Asia/Karachi (PKT)",
                style_meta_right,
            ),
        ]
    ]
    header_table = Table(header_table_data, colWidths=[450, 320])
    header_table.setStyle(
        TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(header_table)
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#E2E8F0"), spaceAfter=8))

    # 2. Summary KPI Ribbon
    completed = sum(1 for r in records if r["status"] == "completed")
    cancelled = sum(1 for r in records if r["status"] == "cancelled")

    if scope in ("manager", "cashier"):
        total_rev = sum(r["total"] for r in records)
        total_tax = sum(r["tax"] for r in records)
        kpi_data = [
            [
                Paragraph(f"<b>{len(records)}</b>", style_kpi_num),
                Paragraph(f"<b>{completed}</b>", style_kpi_num),
                Paragraph(f"<b>{cancelled}</b>", style_kpi_num),
                Paragraph(f"<b>{format_money_str(total_tax)}</b>", style_kpi_num),
                Paragraph(f"<b>{format_money_str(total_rev)}</b>", style_kpi_num),
            ],
            [
                Paragraph("TOTAL ORDERS", style_kpi_lbl),
                Paragraph("COMPLETED", style_kpi_lbl),
                Paragraph("CANCELLED", style_kpi_lbl),
                Paragraph("TAX COLLECTED", style_kpi_lbl),
                Paragraph("NET REVENUE", style_kpi_lbl),
            ],
        ]
        kpi_cols = [154, 154, 154, 154, 154]
    elif scope == "kitchen":
        durations = [r["prep_minutes"] for r in records if r.get("prep_minutes") is not None]
        avg_prep = round(sum(durations) / len(durations), 1) if durations else 0
        total_dishes = sum(sum(it["quantity"] for it in r["items"]) for r in records)
        kpi_data = [
            [
                Paragraph(f"<b>{len(records)}</b>", style_kpi_num),
                Paragraph(f"<b>{completed}</b>", style_kpi_num),
                Paragraph(f"<b>{cancelled}</b>", style_kpi_num),
                Paragraph(f"<b>{total_dishes}</b>", style_kpi_num),
                Paragraph(f"<b>{avg_prep} min</b>", style_kpi_num),
            ],
            [
                Paragraph("TOTAL TICKETS", style_kpi_lbl),
                Paragraph("SERVED / PLATED", style_kpi_lbl),
                Paragraph("CANCELLED", style_kpi_lbl),
                Paragraph("DISHES PREPARED", style_kpi_lbl),
                Paragraph("AVG PREP TIME", style_kpi_lbl),
            ],
        ]
        kpi_cols = [154, 154, 154, 154, 154]
    else:  # waiter
        total_sales = sum(r["total"] for r in records)
        kpi_data = [
            [
                Paragraph(f"<b>{len(records)}</b>", style_kpi_num),
                Paragraph(f"<b>{completed}</b>", style_kpi_num),
                Paragraph(f"<b>{cancelled}</b>", style_kpi_num),
                Paragraph(f"<b>{format_money_str(total_sales)}</b>", style_kpi_num),
            ],
            [
                Paragraph("MY ORDERS", style_kpi_lbl),
                Paragraph("COMPLETED", style_kpi_lbl),
                Paragraph("CANCELLED", style_kpi_lbl),
                Paragraph("RECORDED TOTAL", style_kpi_lbl),
            ],
        ]
        kpi_cols = [192, 192, 192, 194]

    kpi_table = Table(kpi_data, colWidths=kpi_cols)
    kpi_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(kpi_table)
    story.append(Spacer(1, 10))

    # 3. Data Table (Landscape with repeatRows=1)
    if scope == "manager":
        table_headers = [
            Paragraph("Order #", style_th),
            Paragraph("Date / Time", style_th),
            Paragraph("Table / Floor", style_th),
            Paragraph("Waiter", style_th),
            Paragraph("Status", style_th),
            Paragraph("Items Ordered", style_th),
            Paragraph("Subtotal", style_th),
            Paragraph("Total (PKR)", style_th),
        ]
        col_widths = [65, 85, 80, 80, 65, 235, 75, 85]

        table_data = [table_headers]
        for r in records:
            status_color = "#16A34A" if r["status"] == "completed" else "#DC2626" if r["status"] == "cancelled" else "#D97706"
            items_p = Paragraph(r["items_summary"] or "—", style_td)
            table_data.append([
                Paragraph(f"<b>{r['order_number']}</b>", style_td),
                Paragraph(format_karachi_dt(r["created_at"]), style_td_center),
                Paragraph(r["floor_table"], style_td),
                Paragraph(r["waiter_name"], style_td),
                Paragraph(f"<font color='{status_color}'><b>{r['status'].upper()}</b></font>", style_td_center),
                items_p,
                Paragraph(format_money_str(r["subtotal"]), style_td_right),
                Paragraph(f"<b>{format_money_str(r['total'])}</b>", style_td_right),
            ])

    elif scope == "cashier":
        table_headers = [
            Paragraph("Receipt #", style_th),
            Paragraph("Order #", style_th),
            Paragraph("Time", style_th),
            Paragraph("Table", style_th),
            Paragraph("Status", style_th),
            Paragraph("Subtotal", style_th),
            Paragraph("Discount", style_th),
            Paragraph("Tax", style_th),
            Paragraph("Net Total (PKR)", style_th),
            Paragraph("Cashier", style_th),
        ]
        col_widths = [75, 65, 85, 75, 60, 75, 65, 65, 85, 75]

        table_data = [table_headers]
        for r in records:
            status_color = "#16A34A" if r["status"] == "completed" else "#DC2626" if r["status"] == "cancelled" else "#D97706"
            table_data.append([
                Paragraph(f"<b>{r['receipt_number']}</b>", style_td),
                Paragraph(r["order_number"], style_td),
                Paragraph(format_karachi_dt(r["created_at"]), style_td_center),
                Paragraph(r["floor_table"], style_td),
                Paragraph(f"<font color='{status_color}'><b>{r['status'].upper()}</b></font>", style_td_center),
                Paragraph(format_money_str(r["subtotal"]), style_td_right),
                Paragraph(format_money_str(r["discount"]), style_td_right),
                Paragraph(format_money_str(r["tax"]), style_td_right),
                Paragraph(f"<b>{format_money_str(r['total'])}</b>", style_td_right),
                Paragraph(r["cashier_name"], style_td),
            ])

    elif scope == "waiter":
        table_headers = [
            Paragraph("Order #", style_th),
            Paragraph("Date / Time", style_th),
            Paragraph("Table / Floor", style_th),
            Paragraph("Status", style_th),
            Paragraph("Items Ordered", style_th),
            Paragraph("Total (PKR)", style_th),
            Paragraph("Special Notes", style_th),
        ]
        col_widths = [75, 95, 95, 70, 245, 90, 100]

        table_data = [table_headers]
        for r in records:
            status_color = "#16A34A" if r["status"] == "completed" else "#DC2626" if r["status"] == "cancelled" else "#D97706"
            table_data.append([
                Paragraph(f"<b>{r['order_number']}</b>", style_td),
                Paragraph(format_karachi_dt(r["created_at"]), style_td_center),
                Paragraph(r["floor_table"], style_td),
                Paragraph(f"<font color='{status_color}'><b>{r['status'].upper()}</b></font>", style_td_center),
                Paragraph(r["items_summary"] or "—", style_td),
                Paragraph(f"<b>{format_money_str(r['total'])}</b>", style_td_right),
                Paragraph(r["notes"] or "—", style_td),
            ])

    else:  # kitchen
        table_headers = [
            Paragraph("Ticket #", style_th),
            Paragraph("Order Time", style_th),
            Paragraph("Duration", style_th),
            Paragraph("Table / Floor", style_th),
            Paragraph("Status", style_th),
            Paragraph("Dishes & Quantities", style_th),
            Paragraph("Chef Dispatcher", style_th),
            Paragraph("Special Notes", style_th),
        ]
        col_widths = [75, 95, 60, 95, 65, 235, 75, 70]

        table_data = [table_headers]
        for r in records:
            status_color = "#16A34A" if r["status"] == "completed" else "#DC2626" if r["status"] == "cancelled" else "#D97706"
            prep_str = f"<b>{r['prep_minutes']} min</b>" if r.get("prep_minutes") is not None else "—"
            table_data.append([
                Paragraph(f"<b>TCK-{r['order_number']}</b>", style_td),
                Paragraph(format_karachi_dt(r["created_at"]), style_td_center),
                Paragraph(prep_str, style_td_center),
                Paragraph(r["floor_table"], style_td),
                Paragraph(f"<font color='{status_color}'><b>{r['status'].upper()}</b></font>", style_td_center),
                Paragraph(r["items_summary"] or "—", style_td),
                Paragraph(r["chef_name"], style_td),
                Paragraph(r["notes"] or "—", style_td),
            ])

    main_table = Table(table_data, colWidths=col_widths, repeatRows=1)

    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]

    for i in range(1, len(table_data)):
        if i % 2 == 0:
            t_style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F8FAFC")))

    main_table.setStyle(TableStyle(t_style))
    story.append(main_table)

    doc.build(story, canvasmaker=NumberedCanvas)
    buf.seek(0)
    return buf
